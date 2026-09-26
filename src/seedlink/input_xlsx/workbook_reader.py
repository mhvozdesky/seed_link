"""Strict, provenance-preserving import of the four supported XLSX exports."""

from __future__ import annotations

from collections import Counter, defaultdict
from collections.abc import Callable, Mapping
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta
from decimal import Decimal
from enum import StrEnum
from hashlib import sha256
from io import BytesIO
import logging
from os import PathLike
from pathlib import Path
from types import MappingProxyType
from typing import TypeAlias
import traceback
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile, is_zipfile

from openpyxl import load_workbook

from seedlink import __version__
from seedlink.domain.issues import (
    ISSUE_LEVEL_BY_CODE,
    Issue,
    IssueCode,
    IssueLevel,
)
from seedlink.domain.provenance import (
    InputRole,
    InputSnapshot,
    InputSource,
    SourceCell,
    SourceRecord,
    SourceValueKind,
)
from seedlink.domain.salesforce_ids import salesforce_id_key
from seedlink.input_xlsx import r1, r2, r3, r4
from seedlink.input_xlsx._parsing import ParseProblem
from seedlink.input_xlsx.schemas import SCHEMAS, SchemaDefinition, schema_for


ParsedRecord: TypeAlias = r1.R1Record | r2.R2Record | r3.R3Record | r4.R4Record

_LOGGER = logging.getLogger("seedlink.input_xlsx")
_LOGGER.addHandler(logging.NullHandler())


class WorkbookContentError(Exception):
    """Raised only when lazy openpyxl reading fails on workbook bytes."""


class RecordGroupStatus(StrEnum):
    UNIQUE = "unique"
    IDENTICAL_ID = "identical_id"
    CONFLICTING_ID = "conflicting_id"
    UNCERTAIN_NO_ID = "uncertain_no_id"


@dataclass(frozen=True, slots=True)
class RecordGroup:
    """How physical source rows contribute to later entity construction."""

    status: RecordGroupStatus
    identifier_field: str
    identifier: str | None
    record_keys: tuple[str, ...]
    representative_record_key: str | None

    def __post_init__(self) -> None:
        if not self.identifier_field:
            raise ValueError("identifier_field must not be blank")
        if not self.record_keys or len(self.record_keys) != len(set(self.record_keys)):
            raise ValueError("record_keys must be non-empty and unique")
        if self.status in {
            RecordGroupStatus.UNIQUE,
            RecordGroupStatus.IDENTICAL_ID,
        }:
            if self.representative_record_key not in self.record_keys:
                raise ValueError("countable record group needs a representative")
        elif self.representative_record_key is not None:
            raise ValueError("uncertain record group cannot choose a representative")

    @property
    def is_countable(self) -> bool:
        return self.representative_record_key is not None


@dataclass(frozen=True, slots=True)
class RecordAccounting:
    """Countable records paired with groups whose unique count is unknown."""

    countable_records: tuple[ParsedRecord, ...]
    uncertain_groups: tuple[RecordGroup, ...]

    @property
    def is_complete(self) -> bool:
        return not self.uncertain_groups


@dataclass(frozen=True, slots=True)
class WorkbookReadResult:
    expected_role: InputRole
    path: Path
    detected_role: InputRole | None
    source: InputSource | None
    records: tuple[ParsedRecord, ...]
    record_groups: tuple[RecordGroup, ...]
    issues: tuple[Issue, ...]

    @property
    def is_accepted(self) -> bool:
        return self.source is not None and not self.fatal_issues

    @property
    def blocking_issues(self) -> tuple[Issue, ...]:
        return tuple(
            issue
            for issue in self.issues
            if issue.level is IssueLevel.IMPORT_BLOCKING
        )

    @property
    def record_accounting(self) -> RecordAccounting:
        representative_keys = {
            group.representative_record_key
            for group in self.record_groups
            if group.representative_record_key is not None
        }
        countable_records = tuple(
            record
            for record in self.records
            if record.source.key in representative_keys
        )
        uncertain_groups = tuple(
            group for group in self.record_groups if not group.is_countable
        )
        return RecordAccounting(countable_records, uncertain_groups)

    @property
    def countable_records(self) -> tuple[ParsedRecord, ...]:
        return self.record_accounting.countable_records

    @property
    def uncertain_groups(self) -> tuple[RecordGroup, ...]:
        return self.record_accounting.uncertain_groups

    @property
    def fatal_issues(self) -> tuple[Issue, ...]:
        return tuple(
            issue
            for issue in self.issues
            if issue.level in {IssueLevel.IMPORT_BLOCKING, IssueLevel.INTERNAL}
        )


@dataclass(frozen=True, slots=True)
class ImportResult:
    snapshot: InputSnapshot | None
    workbooks: tuple[WorkbookReadResult, ...]
    issues: tuple[Issue, ...]

    @property
    def is_accepted(self) -> bool:
        return self.snapshot is not None

    @property
    def blocking_issues(self) -> tuple[Issue, ...]:
        return tuple(
            issue
            for issue in self.issues
            if issue.level is IssueLevel.IMPORT_BLOCKING
        )

    @property
    def fatal_issues(self) -> tuple[Issue, ...]:
        return tuple(
            issue
            for issue in self.issues
            if issue.level in {IssueLevel.IMPORT_BLOCKING, IssueLevel.INTERNAL}
        )

    def workbook(self, role: InputRole) -> WorkbookReadResult:
        for workbook in self.workbooks:
            if workbook.expected_role is role:
                return workbook
        raise KeyError(role)


_MAIN_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
)
_CONTENT_TYPES_NAMESPACE = {
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types"
}
_IDENTITY_FIELDS = MappingProxyType(
    {
        InputRole.R1: "Record Nr",
        InputRole.R2: "Opportunity Product Id 18",
        InputRole.R3: "Campaign Member Id 18",
        InputRole.R4: "Activity ID",
    }
)
_IDENTITY_ATTRIBUTES = MappingProxyType(
    {
        InputRole.R1: "record_number",
        InputRole.R2: "product_id",
        InputRole.R3: "campaign_member_id",
        InputRole.R4: "activity_id",
    }
)
_PARSERS = MappingProxyType(
    {
        InputRole.R1: r1.parse_record,
        InputRole.R2: r2.parse_record,
        InputRole.R3: r3.parse_record,
        InputRole.R4: r4.parse_record,
    }
)


def _issue_id(
    code: IssueCode,
    role: InputRole,
    *,
    file_sha256: str | None = None,
    sources: tuple[SourceCell, ...] = (),
    discriminator: str = "",
) -> str:
    identity = "\x1f".join(
        (
            code.value,
            role.value,
            file_sha256 or "",
            *(source.stable_key for source in sources),
            discriminator,
        )
    )
    return f"issue:{sha256(identity.encode('utf-8')).hexdigest()}"


def _make_issue(
    code: IssueCode,
    role: InputRole,
    message_uk: str,
    *,
    file_sha256: str | None = None,
    sources: tuple[SourceCell, ...] = (),
    affected_keys: tuple[str, ...] = (),
    details: tuple[tuple[str, str], ...] = (),
    discriminator: str = "",
) -> Issue:
    return Issue(
        issue_id=_issue_id(
            code,
            role,
            file_sha256=file_sha256,
            sources=sources,
            discriminator=discriminator,
        ),
        code=code,
        level=ISSUE_LEVEL_BY_CODE[code],
        message_uk=message_uk,
        sources=sources,
        affected_keys=affected_keys,
        details=details,
    )


def _blocking_result(
    role: InputRole,
    path: Path,
    code: IssueCode,
    message: str,
    *,
    file_sha256: str | None = None,
    detected_role: InputRole | None = None,
    details: tuple[tuple[str, str], ...] = (),
) -> WorkbookReadResult:
    issue = _make_issue(
        code,
        role,
        message,
        file_sha256=file_sha256,
        details=details,
        discriminator=str(path),
    )
    return WorkbookReadResult(role, path, detected_role, None, (), (), (issue,))


def _internal_result(
    role: InputRole,
    path: Path,
    file_sha256: str,
    phase: str,
    error: Exception,
    *,
    detected_role: InputRole | None = None,
    source: InputSource | None = None,
    prior_issues: tuple[Issue, ...] = (),
) -> WorkbookReadResult:
    diagnostic_identity = "\x1f".join(
        (role.value, file_sha256, phase, type(error).__name__)
    )
    diagnostic_ref = sha256(diagnostic_identity.encode("utf-8")).hexdigest()[:12]
    stack = "".join(traceback.format_tb(error.__traceback__))
    _LOGGER.error(
        "Внутрішня помилка імпорту %s; diagnostic_ref=%s; "
        "phase=%s; error_type=%s\nСтек викликів без значень джерела:\n%s",
        role.value,
        diagnostic_ref,
        phase,
        type(error).__name__,
        stack,
    )
    issue = _make_issue(
        IssueCode.INTERNAL_ERROR,
        role,
        f"Внутрішня помилка під час імпорту {role.value}. "
        f"Код діагностики: {diagnostic_ref}.",
        file_sha256=file_sha256,
        details=(
            ("phase", phase),
            ("error_type", type(error).__name__),
            ("diagnostic_ref", diagnostic_ref),
        ),
        discriminator=f"{phase}:{type(error).__name__}",
    )
    return WorkbookReadResult(
        expected_role=role,
        path=path,
        detected_role=detected_role,
        source=source,
        records=(),
        record_groups=(),
        issues=prior_issues + (issue,),
    )


def _validate_ooxml_package(data: bytes) -> str | None:
    """Return an error reason when bytes are not a plain XLSX workbook package."""

    stream = BytesIO(data)
    if not is_zipfile(stream):
        return "файл не є ZIP-пакетом OOXML"
    try:
        with ZipFile(BytesIO(data)) as archive:
            names = set(archive.namelist())
            required = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"}
            missing = sorted(required - names)
            if missing:
                return (
                    "в OOXML-пакеті відсутні обов'язкові частини"
                )
            if any(info.flag_bits & 0x1 for info in archive.infolist()):
                return "зашифровані OOXML-частини не підтримуються"
            if "xl/vbaProject.bin" in names:
                return "книги з макросами не підтримуються"
            try:
                root = ElementTree.fromstring(archive.read("[Content_Types].xml"))
            except (ElementTree.ParseError, KeyError):
                return "пошкоджений опис OOXML-пакета"
            workbook_types = {
                element.attrib.get("ContentType")
                for element in root.findall("ct:Override", _CONTENT_TYPES_NAMESPACE)
                if element.attrib.get("PartName") == "/xl/workbook.xml"
            }
            if workbook_types != {_MAIN_CONTENT_TYPE}:
                return "тип книги не є звичайним XLSX"
    except (BadZipFile, OSError):
        return "пошкоджений ZIP-пакет OOXML"
    return None


def _header_values(worksheet) -> tuple[object, ...]:
    max_column = max(worksheet.max_column or 1, 1)
    try:
        rows = worksheet.iter_rows(
            min_row=1, max_row=1, min_col=1, max_col=max_column
        )
        row = next(rows, ())
    except Exception as exc:
        raise WorkbookContentError("cannot read the header row") from exc
    return tuple(cell.value for cell in row)


def _header_text(value: object) -> str | None:
    return value if isinstance(value, str) and value.strip() else None


def _detect_role(headers: tuple[object, ...]) -> InputRole | None:
    text_headers = tuple(_header_text(value) for value in headers)
    if any(value is None for value in text_headers):
        return None
    if len(text_headers) != len(set(text_headers)):
        return None
    header_set = set(text_headers)
    for role, schema in SCHEMAS.items():
        if len(text_headers) == len(schema.columns) and header_set == set(schema.columns):
            return role
    return None


def _schema_details(
    role: InputRole, headers: tuple[object, ...]
) -> tuple[tuple[str, str], ...]:
    expected = schema_for(role).columns
    valid_text = [value for value in headers if isinstance(value, str) and value.strip()]
    counts = Counter(valid_text)
    duplicates = sorted(value for value, count in counts.items() if count > 1)
    missing = sorted(set(expected) - set(valid_text))
    added = sorted(set(valid_text) - set(expected))
    blanks = [
        str(index)
        for index, value in enumerate(headers, start=1)
        if not isinstance(value, str) or not value.strip()
    ]
    return (
        ("role", role.value),
        ("missing", "; ".join(missing) or "—"),
        ("unexpected", "; ".join(added) or "—"),
        ("duplicate", "; ".join(duplicates) or "—"),
        ("blank_columns", "; ".join(blanks) or "—"),
    )


def _schema_message(role: InputRole, details: tuple[tuple[str, str], ...]) -> str:
    values = dict(details)
    parts: list[str] = []
    if values["missing"] != "—":
        parts.append(f"відсутні: {values['missing']}")
    if values["unexpected"] != "—":
        parts.append(f"додані: {values['unexpected']}")
    if values["duplicate"] != "—":
        parts.append(f"дубльовані: {values['duplicate']}")
    if values["blank_columns"] != "—":
        parts.append(
            f"порожні заголовки у колонках: {values['blank_columns']}"
        )
    difference = "; ".join(parts) or "набір заголовків не збігається"
    return f"Структура {role.value} не підтримується: {difference}."


def _source_value_kind(cell) -> SourceValueKind:
    if cell.data_type == "f":
        return SourceValueKind.FORMULA
    if cell.data_type == "e":
        return SourceValueKind.ERROR
    value = cell.value
    if value is None:
        return SourceValueKind.BLANK
    if isinstance(value, bool):
        return SourceValueKind.BOOLEAN
    if isinstance(value, datetime):
        return SourceValueKind.DATETIME
    if isinstance(value, date):
        return SourceValueKind.DATE
    if isinstance(value, time):
        return SourceValueKind.TIME
    if isinstance(value, timedelta):
        return SourceValueKind.DURATION
    if isinstance(value, str):
        return SourceValueKind.TEXT
    if isinstance(value, (int, float, Decimal)):
        return SourceValueKind.NUMBER
    return SourceValueKind.UNKNOWN


def _record_key(
    role: InputRole, file_sha256: str, sheet_name: str, excel_row: int
) -> str:
    identity = "\x1f".join(
        (role.value, file_sha256, sheet_name, str(excel_row))
    )
    return f"record:{sha256(identity.encode('utf-8')).hexdigest()}"


def _read_records(
    worksheet,
    *,
    path: Path,
    file_sha256: str,
    role: InputRole,
    schema: SchemaDefinition,
    headers: tuple[object, ...],
) -> tuple[tuple[SourceRecord, ...], tuple[Issue, ...]]:
    column_by_name = {value: index for index, value in enumerate(headers, start=1)}
    records: list[SourceRecord] = []
    issues: list[Issue] = []
    max_column = len(headers)
    try:
        worksheet_rows = iter(
            worksheet.iter_rows(min_row=2, min_col=1, max_col=max_column)
        )
    except Exception as exc:
        raise WorkbookContentError("cannot start reading data rows") from exc
    excel_row = 2
    while True:
        try:
            worksheet_row = next(worksheet_rows)
        except StopIteration:
            break
        except Exception as exc:
            raise WorkbookContentError("cannot read a data row") from exc
        cells_by_column = {
            column: cell for column, cell in enumerate(worksheet_row, start=1)
        }
        source_cells: list[SourceCell] = []
        has_value = False
        for field_name in schema.columns:
            cell = cells_by_column[column_by_name[field_name]]
            kind = _source_value_kind(cell)
            if kind is not SourceValueKind.BLANK:
                has_value = True
            hyperlink = getattr(cell, "hyperlink", None)
            hyperlink_target = None
            if hyperlink:
                hyperlink_target = (
                    getattr(hyperlink, "target", None)
                    or getattr(hyperlink, "location", None)
                    or None
                )
            source_cells.append(
                SourceCell(
                    role=role,
                    file_name=path.name,
                    file_sha256=file_sha256,
                    sheet_name=worksheet.title,
                    excel_row=excel_row,
                    field_name=field_name,
                    original_value=cell.value,
                    value_kind=kind,
                    hyperlink_target=hyperlink_target,
                )
            )
        if not has_value:
            excel_row += 1
            continue
        record = SourceRecord(
            key=_record_key(role, file_sha256, worksheet.title, excel_row),
            role=role,
            file_name=path.name,
            file_sha256=file_sha256,
            sheet_name=worksheet.title,
            excel_row=excel_row,
            cells=tuple(source_cells),
        )
        records.append(record)
        for source_cell in source_cells:
            if source_cell.value_kind not in {
                SourceValueKind.FORMULA,
                SourceValueKind.ERROR,
                SourceValueKind.UNKNOWN,
            }:
                continue
            kind_uk = {
                SourceValueKind.FORMULA: "формулу",
                SourceValueKind.ERROR: "помилку Excel",
                SourceValueKind.UNKNOWN: "непідтримуваний тип",
            }[source_cell.value_kind]
            issues.append(
                _make_issue(
                    IssueCode.UNSUPPORTED_CELL_VALUE,
                    role,
                    f"Поле «{source_cell.field_name}» містить {kind_uk}; "
                    "значення не використано.",
                    file_sha256=file_sha256,
                    sources=(source_cell,),
                    affected_keys=(record.key,),
                )
            )
        excel_row += 1
    return tuple(records), tuple(issues)


def _problem_issue(
    problem: ParseProblem, role: InputRole, record_key: str, file_sha256: str
) -> Issue:
    return _make_issue(
        problem.code,
        role,
        problem.message_uk,
        file_sha256=file_sha256,
        sources=(problem.cell,),
        affected_keys=(record_key,),
    )


def _record_signature(record: SourceRecord) -> tuple[tuple[str, str, str, str], ...]:
    return tuple(
        (
            cell.field_name,
            cell.value_kind.value if cell.value_kind else "unknown",
            repr(cell.original_value),
            cell.hyperlink_target or "",
        )
        for cell in record.cells
    )


def _different_fields(records: tuple[SourceRecord, ...]) -> tuple[str, ...]:
    result: list[str] = []
    for field in (cell.field_name for cell in records[0].cells):
        values = {
            (
                repr(record.cell(field).original_value),
                record.cell(field).hyperlink_target,
                record.cell(field).value_kind,
            )
            for record in records
        }
        if len(values) > 1:
            result.append(field)
    return tuple(result)


def _group_records(
    role: InputRole,
    parsed_records: tuple[ParsedRecord, ...],
    file_sha256: str,
) -> tuple[tuple[RecordGroup, ...], tuple[Issue, ...]]:
    identity_field = _IDENTITY_FIELDS[role]
    identity_attribute = _IDENTITY_ATTRIBUTES[role]
    by_id: dict[str, list[ParsedRecord]] = defaultdict(list)
    without_id: list[ParsedRecord] = []
    for record in parsed_records:
        identifier = getattr(record, identity_attribute)
        if role is InputRole.R3 and identifier is not None:
            identifier = salesforce_id_key(identifier)
        (by_id[identifier] if identifier is not None else without_id).append(record)

    groups: list[RecordGroup] = []
    issues: list[Issue] = []
    for identifier, records_with_id in by_id.items():
        records = tuple(record.source for record in records_with_id)
        keys = tuple(record.key for record in records)
        if len(records) == 1:
            groups.append(
                RecordGroup(
                    RecordGroupStatus.UNIQUE,
                    identity_field,
                    identifier,
                    keys,
                    keys[0],
                )
            )
            continue
        sources = tuple(record.cell(identity_field) for record in records)
        signatures = {_record_signature(record) for record in records}
        if len(signatures) == 1:
            groups.append(
                RecordGroup(
                    RecordGroupStatus.IDENTICAL_ID,
                    identity_field,
                    identifier,
                    keys,
                    keys[0],
                )
            )
            issues.append(
                _make_issue(
                    IssueCode.DUPLICATE_ID,
                    role,
                    f"Ідентичні записи з {identity_field}={identifier} "
                    "враховано один раз.",
                    file_sha256=file_sha256,
                    sources=sources,
                    affected_keys=keys,
                    discriminator=identifier,
                )
            )
        else:
            different = _different_fields(records)
            conflict_code = (
                IssueCode.PRODUCT_ID_CONFLICT
                if role is InputRole.R2
                else IssueCode.CONFLICTING_ID_DATA
            )
            detail_sources = tuple(
                record.cell(field)
                for field in different
                for record in records
            )
            groups.append(
                RecordGroup(
                    RecordGroupStatus.CONFLICTING_ID,
                    identity_field,
                    identifier,
                    keys,
                    None,
                )
            )
            issues.append(
                _make_issue(
                    conflict_code,
                    role,
                    f"Записи з {identity_field}={identifier} "
                    "суперечать один одному.",
                    file_sha256=file_sha256,
                    sources=sources + detail_sources,
                    affected_keys=keys,
                    details=(("different_fields", "; ".join(different)),),
                    discriminator=identifier,
                )
            )

    by_signature: dict[
        tuple[tuple[str, str, str, str], ...], list[ParsedRecord]
    ] = defaultdict(list)
    for record in without_id:
        by_signature[_record_signature(record.source)].append(record)
    for signature_records in by_signature.values():
        records = tuple(record.source for record in signature_records)
        keys = tuple(record.key for record in records)
        if len(records) == 1:
            groups.append(
                RecordGroup(
                    RecordGroupStatus.UNIQUE,
                    identity_field,
                    None,
                    keys,
                    keys[0],
                )
            )
            continue
        sources = tuple(record.cell(identity_field) for record in records)
        groups.append(
            RecordGroup(
                RecordGroupStatus.UNCERTAIN_NO_ID,
                identity_field,
                None,
                keys,
                None,
            )
        )
        issues.append(
            _make_issue(
                IssueCode.DUPLICATE_UNCERTAIN,
                role,
                f"Однакові записи без {identity_field} неможливо "
                "надійно відрізнити.",
                file_sha256=file_sha256,
                sources=sources,
                affected_keys=keys,
                discriminator="|".join(keys),
            )
        )

    excel_row_by_key = {
        record.source.key: record.source.excel_row for record in parsed_records
    }
    groups.sort(
        key=lambda group: min(excel_row_by_key[key] for key in group.record_keys)
    )
    return tuple(groups), tuple(issues)


def read_workbook(
    path: str | PathLike[str], expected_role: InputRole
) -> WorkbookReadResult:
    """Read one slot, validating extension, OOXML package, structure and schema."""

    source_path = Path(path)
    if not source_path.exists() or not source_path.is_file():
        return _blocking_result(
            expected_role,
            source_path,
            IssueCode.MISSING_INPUT,
            f"Для блоку {expected_role.label_uk} файл не знайдено.",
        )
    if source_path.suffix.casefold() != ".xlsx":
        return _blocking_result(
            expected_role,
            source_path,
            IssueCode.INVALID_XLSX,
            f"Блок {expected_role.label_uk} приймає лише файл .xlsx.",
        )
    try:
        data = source_path.read_bytes()
    except OSError as exc:
        return _blocking_result(
            expected_role,
            source_path,
            IssueCode.WORKBOOK_UNREADABLE,
            f"Не вдалося прочитати файл для "
            f"{expected_role.label_uk}: {exc}.",
        )
    file_sha256 = sha256(data).hexdigest()
    invalid_reason = _validate_ooxml_package(data)
    if invalid_reason:
        return _blocking_result(
            expected_role,
            source_path,
            IssueCode.INVALID_XLSX,
            f"Файл для {expected_role.label_uk} не є підтримуваною "
            f"книгою XLSX: {invalid_reason}.",
            file_sha256=file_sha256,
        )

    workbook = None
    detected_role: InputRole | None = None
    source: InputSource | None = None
    prior_issues: tuple[Issue, ...] = ()
    phase = "openpyxl_load"
    try:
        try:
            workbook = load_workbook(
                BytesIO(data),
                read_only=expected_role is not InputRole.R1,
                data_only=False,
                keep_links=False,
            )
        except Exception as exc:
            return _blocking_result(
                expected_role,
                source_path,
                IssueCode.INVALID_XLSX,
                f"Файл для {expected_role.label_uk} пошкоджений або має "
                f"непідтримувану структуру: {type(exc).__name__}.",
                file_sha256=file_sha256,
            )

        phase = "structure_validation"
        if len(workbook.sheetnames) != 1 or len(workbook.worksheets) != 1:
            return _blocking_result(
                expected_role,
                source_path,
                IssueCode.UNSUPPORTED_STRUCTURE,
                f"Файл для {expected_role.label_uk} має містити "
                "рівно один аркуш даних.",
                file_sha256=file_sha256,
                details=(("sheet_count", str(len(workbook.sheetnames))),),
            )
        worksheet = workbook.worksheets[0]
        headers = _header_values(worksheet)
        detected_role = _detect_role(headers)
        if detected_role is not None and detected_role is not expected_role:
            return _blocking_result(
                expected_role,
                source_path,
                IssueCode.ROLE_MISMATCH,
                f"У блоці {expected_role.value} вибрано файл структури "
                f"{detected_role.value}. Виберіть правильний експорт "
                f"для «{expected_role.label_uk}».",
                file_sha256=file_sha256,
                detected_role=detected_role,
                details=(
                    ("expected_role", expected_role.value),
                    ("detected_role", detected_role.value),
                ),
            )
        schema = schema_for(expected_role)
        if detected_role is not expected_role:
            details = _schema_details(expected_role, headers)
            return _blocking_result(
                expected_role,
                source_path,
                IssueCode.SCHEMA_MISMATCH,
                _schema_message(expected_role, details),
                file_sha256=file_sha256,
                details=details,
            )

        phase = "source_rows"
        raw_records, cell_issues = _read_records(
            worksheet,
            path=source_path,
            file_sha256=file_sha256,
            role=expected_role,
            schema=schema,
            headers=headers,
        )
        prior_issues = cell_issues
        source = InputSource(
            role=expected_role,
            file_name=source_path.name,
            file_sha256=file_sha256,
            sheet_name=worksheet.title,
            schema_version=schema.version,
            row_count=len(raw_records),
            records=raw_records,
        )

        phase = "role_parsing"
        parsed_records: list[ParsedRecord] = []
        parse_issues: list[Issue] = []
        parser = _PARSERS[expected_role]
        for record in raw_records:
            parsed, problems = parser(record)
            parsed_records.append(parsed)
            parse_issues.extend(
                _problem_issue(problem, expected_role, record.key, file_sha256)
                for problem in problems
            )
        prior_issues = cell_issues + tuple(parse_issues)
        phase = "record_grouping"
        record_groups, duplicate_issues = _group_records(
            expected_role, tuple(parsed_records), file_sha256
        )
        return WorkbookReadResult(
            expected_role=expected_role,
            path=source_path,
            detected_role=expected_role,
            source=source,
            records=tuple(parsed_records),
            record_groups=record_groups,
            issues=cell_issues + tuple(parse_issues) + duplicate_issues,
        )
    except WorkbookContentError as exc:
        return _blocking_result(
            expected_role,
            source_path,
            IssueCode.INVALID_XLSX,
            f"Не вдалося прочитати вміст книги для "
            f"{expected_role.label_uk}: пошкоджений OOXML.",
            file_sha256=file_sha256,
        )
    except Exception as exc:
        return _internal_result(
            expected_role,
            source_path,
            file_sha256,
            phase,
            exc,
            detected_role=detected_role,
            source=source,
            prior_issues=prior_issues,
        )
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                _LOGGER.warning(
                    "Не вдалося закрити in-memory книгу для %s.",
                    expected_role.value,
                )


def _snapshot_id(sources: tuple[InputSource, ...]) -> str:
    identity = "\x1e".join(
        f"{source.role.value}\x1f{source.file_sha256}\x1f{source.schema_version}"
        for source in sources
    )
    return f"snapshot:{sha256(identity.encode('utf-8')).hexdigest()}"


def import_workbooks(
    paths: Mapping[InputRole, str | PathLike[str] | None],
    *,
    loaded_at: datetime | None = None,
    program_version: str = __version__,
    check_cancelled: Callable[[], None] | None = None,
    workbook_loaded: Callable[[InputRole, int, int], None] | None = None,
) -> ImportResult:
    """Read all slots and build a snapshot only when no fatal issue exists."""

    checkpoint = check_cancelled or (lambda: None)
    results: list[WorkbookReadResult] = []
    roles = tuple(InputRole)
    checkpoint()
    for index, role in enumerate(roles, start=1):
        path = paths.get(role)
        if path is None:
            results.append(
                _blocking_result(
                    role,
                    Path(f"<{role.value} не вибрано>"),
                    IssueCode.MISSING_INPUT,
                    f"Для блоку {role.label_uk} не вибрано файл.",
                )
            )
        else:
            results.append(read_workbook(path, role))
        if workbook_loaded is not None:
            workbook_loaded(role, index, len(roles))
        checkpoint()
    all_issues = tuple(issue for result in results for issue in result.issues)
    if any(not result.is_accepted for result in results):
        return ImportResult(None, tuple(results), all_issues)

    sources = tuple(
        result.source for result in results if result.source is not None
    )
    timestamp = loaded_at or datetime.now(UTC)
    snapshot = InputSnapshot(
        snapshot_id=_snapshot_id(sources),
        program_version=program_version,
        loaded_at=timestamp,
        sources=sources,
    )
    snapshot.require_complete()
    return ImportResult(snapshot, tuple(results), all_issues)
