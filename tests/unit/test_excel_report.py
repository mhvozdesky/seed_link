from __future__ import annotations

from datetime import UTC, date, datetime
from zipfile import ZipFile

from openpyxl import load_workbook

from seedlink.application import analyze_links, build_report_result
from seedlink.domain import InputRole
from seedlink.input_xlsx import import_workbooks
from seedlink.reports import write_excel_report
from seedlink.reports.excel import MAIN_SHEETS
from seedlink.reports.excel_formulas import (
    completeness_formula,
    selected_by_visible_key_formula,
    visible_unique_count_formula,
    visibility_formula,
)


MEMBER = "00vAbCdEfGhIjKlIVK"


def _result(workbook_set_factory):
    paths = workbook_set_factory(
        {
            InputRole.R1: [
                {
                    "Record Nr": "R1",
                    "Survey Response ID": "S1",
                    "Campaigm Member": MEMBER,
                    "Time Taken": "2026-09-01",
                    "Answer": "=1+1",
                    "Answer (Long Text)": "SE-100",
                }
            ],
            InputRole.R2: [
                {
                    "Opportunity Product Id 18": "P1",
                    "Voucher Number": "SE-100",
                    "Current Year Planned Quantity": "0.1",
                    "Tax ID 1": "00000001",
                    "Account Name": "Клієнт А",
                    "Species group": "SUN: Sunflowers",
                    "Parent Product Local Description": "Гібрид А",
                    "Opportunity Product: Created Date": date(2026, 9, 2),
                },
                {
                    "Opportunity Product Id 18": "P2",
                    "Voucher Number": "SE-100",
                    "Current Year Planned Quantity": "0.2",
                    "Tax ID 1": "00000002",
                    "Account Name": "Клієнт Б",
                    "Species group": "CRN: Corn Seeds",
                    "Parent Product Local Description": None,
                    "Opportunity Product: Created Date": date(2026, 9, 3),
                },
                {
                    "Opportunity Product Id 18": "P3",
                    "Voucher Number": "SE-200",
                    "Current Year Planned Quantity": 4,
                    "Tax ID 1": "00000001",
                    "Account Name": "Клієнт А",
                    "Species group": "SUN: Sunflowers",
                    "Parent Product Local Description": "Гібрид Б",
                    "Opportunity Product: Created Date": date(2026, 8, 31),
                },
            ],
            InputRole.R3: [
                {
                    "Campaign Member Id 18": MEMBER,
                    "First Name": "Тест",
                    "Last Name": "Лід",
                    "Member Type": "Lead",
                    "Member First Associated Date": date(2026, 9, 1),
                }
            ],
            InputRole.R4: [
                {"Activity ID": "A1", "Name": "Тест Лід", "Results": "Результат"}
            ],
        }
    )
    imported = import_workbooks(paths)
    assert imported.is_accepted
    return build_report_result(
        analyze_links(imported), calculated_at=datetime(2026, 9, 10, tzinfo=UTC)
    )


def _row_for_label(sheet, label: str) -> int:
    return next(cell.row for cell in sheet["A"] if cell.value == label)


def test_excel_formula_builders_use_visible_key_sets() -> None:
    assert visibility_formula() == "=SUBTOTAL(103,[@[Ключ]])"
    unique = visible_unique_count_formula("tblFacts", "Voucher")
    assert "FILTER(" in unique and "UNIQUE(" in unique and "Видимий" in unique
    selected = selected_by_visible_key_formula("tblPairs", "Voucher", "Voucher")
    assert "COUNTIFS(" in selected and "Видимий" in selected
    status = completeness_formula("B5", "C5", "E5")
    assert all(label in status for label in ("немає даних", "повний", "частковий", "недоступний"))


def test_excel_report_structure_formulas_caches_and_literal_source_text(
    tmp_path, workbook_set_factory
) -> None:
    result = _result(workbook_set_factory)
    output = tmp_path / "seedlink.xlsx"
    returned = write_excel_report(
        result, output, exported_at=datetime(2026, 9, 11, tzinfo=UTC)
    )
    assert returned == output

    formulas = load_workbook(output, data_only=False, read_only=False)
    assert formulas.sheetnames[:6] == list(MAIN_SHEETS)
    assert formulas["_Факти"].sheet_state == "hidden"
    assert len(formulas["Огляд"]._charts) == 4
    assert "tblLeadVoucher" in formulas["Ліди–ваучери"].tables
    assert "tblMainProducts" in formulas["Товарні рядки"].tables
    assert "tblParticipants" in formulas["Учасники та відповіді"].tables
    assert "tblOtherVouchers" in formulas["Інші ваучери клієнтів"].tables
    assert "tblIssues" in formulas["Проблеми"].tables

    all_formulas = [
        getattr(cell.value, "text", cell.value)
        for sheet in formulas.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]
    assert any("SUBTOTAL" in value for value in all_formulas)
    assert any("FILTER" in value and "UNIQUE" in value for value in all_formulas)
    literal = next(
        cell
        for row in formulas["Джерела"].iter_rows()
        for cell in row
        if cell.value == "=1+1"
    )
    assert literal.data_type == "s"
    assert formulas.calculation.fullCalcOnLoad
    formulas.close()

    cached = load_workbook(output, data_only=True, read_only=False)
    lead_sheet = cached["Ліди–ваучери"]
    quantity_row = _row_for_label(lead_sheet, "Обсяг унікальних видимих ваучерів")
    assert abs(lead_sheet.cell(quantity_row, 2).value - 0.3) < 1e-12
    assert lead_sheet["B6"].value == 1
    assert lead_sheet["B7"].value == 2
    assert lead_sheet.cell(quantity_row, 4).value == "повний"
    assert lead_sheet.cell(quantity_row, 5).value == 0
    other_sheet = cached["Інші ваучери клієнтів"]
    other_row = _row_for_label(other_sheet, "Видимий обсяг")
    assert other_sheet.cell(other_row, 2).value == 4
    cached.close()

    with ZipFile(output) as archive:
        names = set(archive.namelist())
        assert "xl/vbaProject.bin" not in names
        assert not any(name.startswith("xl/externalLinks/") for name in names)


def test_excel_report_accepts_empty_business_tables(tmp_path, workbook_set_factory) -> None:
    imported = import_workbooks(workbook_set_factory({}))
    assert imported.is_accepted
    result = build_report_result(analyze_links(imported))
    output = write_excel_report(result, tmp_path / "empty.xlsx")
    workbook = load_workbook(output, data_only=True)
    assert _row_for_label(workbook["Ліди–ваучери"], "Видимі пари")
    assert workbook["Ліди–ваучери"]["B5"].value == 0
    assert workbook["Товарні рядки"]["D8"].value == "немає даних"
    workbook.close()
