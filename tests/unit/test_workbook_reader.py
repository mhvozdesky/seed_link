from __future__ import annotations

from datetime import datetime, UTC
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook

from seedlink.domain.issues import IssueCode
from seedlink.domain.provenance import InputRole, SourceValueKind
from seedlink.input_xlsx import (
    RecordGroupStatus,
    import_workbooks,
    read_workbook,
    salesforce_id_18,
    salesforce_id_key,
)
from seedlink.input_xlsx import workbook_reader
from seedlink.input_xlsx.schemas import schema_for


def _row(role: InputRole, **values: object) -> dict[str, object]:
    return {column: values.get(column) for column in schema_for(role).columns}


def _write_workbook(
    path: Path,
    role: InputRole,
    rows: list[dict[str, object]],
    *,
    headers: list[object] | None = None,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Дані"
    actual_headers = headers or list(schema_for(role).columns)
    worksheet.append(actual_headers)
    for values in rows:
        worksheet.append([values.get(header) for header in actual_headers])
    workbook.save(path)


def _empty_set(tmp_path: Path) -> dict[InputRole, Path]:
    paths: dict[InputRole, Path] = {}
    for role in InputRole:
        path = tmp_path / f"{role.value}.xlsx"
        _write_workbook(path, role, [])
        paths[role] = path
    return paths


def test_imports_empty_valid_workbooks_and_builds_stable_metadata(tmp_path):
    paths = _empty_set(tmp_path)
    before = {role: path.read_bytes() for role, path in paths.items()}
    loaded_at = datetime(2026, 9, 20, 12, 0, tzinfo=UTC)

    first = import_workbooks(paths, loaded_at=loaded_at, program_version="test")
    second = import_workbooks(paths, loaded_at=loaded_at, program_version="test")

    assert first.is_accepted
    assert first.snapshot is not None
    assert first.snapshot.snapshot_id == second.snapshot.snapshot_id
    assert first.snapshot.loaded_at == loaded_at
    assert {source.row_count for source in first.snapshot.sources} == {0}
    assert not hasattr(first.snapshot.source(InputRole.R1), "source_path")
    assert all(
        source.file_sha256 == sha256(before[source.role]).hexdigest()
        for source in first.snapshot.sources
    )
    assert {role: path.read_bytes() for role, path in paths.items()} == before


def test_column_order_is_irrelevant_and_provenance_uses_excel_row(tmp_path):
    role = InputRole.R4
    headers = list(reversed(schema_for(role).columns))
    path = tmp_path / "довільна назва.xlsx"
    _write_workbook(
        path,
        role,
        [
            _row(
                role,
                **{
                    "Activity ID": "00T000000000001",
                    "Name": "  Учасник   Один  ",
                    "Completed Date/Time": datetime(2026, 9, 20, 10, 30),
                },
            )
        ],
        headers=headers,
    )

    result = read_workbook(path, role)

    assert result.is_accepted
    assert result.source is not None
    assert result.source.sheet_name == "Дані"
    assert result.source.row_count == 1
    record = result.records[0]
    assert record.activity_id == "00T000000000001"
    assert record.person_name == "Учасник Один"
    assert record.source.excel_row == 2
    assert record.source.cell("Activity ID").original_value == "00T000000000001"


def test_wrong_slot_is_reported_as_role_mismatch(tmp_path):
    path = tmp_path / "anything.xlsx"
    _write_workbook(path, InputRole.R2, [])

    result = read_workbook(path, InputRole.R1)

    assert not result.is_accepted
    assert result.detected_role is InputRole.R2
    assert result.issues[0].code is IssueCode.ROLE_MISMATCH
    assert dict(result.issues[0].details) == {
        "expected_role": "R1",
        "detected_role": "R2",
    }


def test_schema_difference_reports_missing_added_duplicate_and_blank(tmp_path):
    headers = list(schema_for(InputRole.R3).columns)
    headers[0] = "Renamed Status"
    headers[1] = "Last Name"
    headers[2] = ""
    path = tmp_path / "schema.xlsx"
    _write_workbook(path, InputRole.R3, [], headers=headers)

    result = read_workbook(path, InputRole.R3)

    assert not result.is_accepted
    issue = result.issues[0]
    assert issue.code is IssueCode.SCHEMA_MISMATCH
    details = dict(issue.details)
    assert "Member Status" in details["missing"]
    assert "Member Type" in details["missing"]
    assert details["unexpected"] == "Renamed Status"
    assert details["duplicate"] == "Last Name"
    assert details["blank_columns"] == "3"


def test_non_xlsx_content_and_extra_sheet_are_blocking(tmp_path):
    fake = tmp_path / "renamed.xlsx"
    fake.write_text("a,b\n1,2\n", encoding="utf-8")
    invalid = read_workbook(fake, InputRole.R1)
    assert invalid.issues[0].code is IssueCode.INVALID_XLSX

    fake_zip = tmp_path / "zip-signature-only.xlsx"
    with ZipFile(fake_zip, "w") as archive:
        archive.writestr("not-a-workbook.txt", "PK is not enough")
    invalid_zip = read_workbook(fake_zip, InputRole.R1)
    assert invalid_zip.issues[0].code is IssueCode.INVALID_XLSX

    multiple = tmp_path / "multiple.xlsx"
    _write_workbook(multiple, InputRole.R1, [])
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(schema_for(InputRole.R1).columns)
    workbook.create_sheet("Другий")
    workbook.save(multiple)
    unsupported = read_workbook(multiple, InputRole.R1)
    assert unsupported.issues[0].code is IssueCode.UNSUPPORTED_STRUCTURE


def test_hidden_rows_are_read_and_formulas_are_not_evaluated(tmp_path):
    path = tmp_path / "formula.xlsx"
    row = _row(
        InputRole.R2,
        **{
            "Opportunity Product Id 18": "00k000000000001",
            "Voucher Number": "SE-001",
            "Current Year Planned Quantity": "=1+2",
        },
    )
    _write_workbook(path, InputRole.R2, [row])
    # Reopen with openpyxl rather than replacing content so row 2 can be hidden.
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    workbook.active.row_dimensions[2].hidden = True
    workbook.save(path)
    workbook.close()

    result = read_workbook(path, InputRole.R2)

    assert result.is_accepted
    assert result.source is not None and result.source.row_count == 1
    quantity_cell = result.source.records[0].cell("Current Year Planned Quantity")
    assert quantity_cell.original_value == "=1+2"
    assert quantity_cell.value_kind is SourceValueKind.FORMULA
    assert result.records[0].quantity is None
    assert [issue.code for issue in result.issues] == [
        IssueCode.UNSUPPORTED_CELL_VALUE
    ]


def test_text_and_native_campaign_member_links_are_inertly_extracted(tmp_path):
    path = tmp_path / "links.xlsx"
    rows = [
        _row(
            InputRole.R1,
            **{
                "Record Nr": "SR-1",
                "Campaigm Member": (
                    '<a href="/lightning/r/CampaignMember/'
                    '00vABCDEF123456XYZ/view" onclick="alert(1)">учасник</a>'
                ),
            },
        ),
        _row(
            InputRole.R1,
            **{"Record Nr": "SR-2", "Campaigm Member": "відкрити"},
        ),
        _row(
            InputRole.R1,
            **{
                "Record Nr": "SR-3",
                "Campaigm Member": (
                    '<a href="/record/701AbCdEfGhIjKl">інший об’єкт</a>'
                ),
            },
        ),
        _row(
            InputRole.R1,
            **{
                "Record Nr": "SR-4",
                "Campaigm Member": '<a href="/record/no-id">зламано</a>',
            },
        ),
    ]
    _write_workbook(path, InputRole.R1, rows)
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    worksheet = workbook.active
    column = list(schema_for(InputRole.R1).columns).index("Campaigm Member") + 1
    worksheet.cell(3, column).hyperlink = (
        "https://example.invalid/00v123456789ABCDEF"
    )
    workbook.save(path)
    workbook.close()

    result = read_workbook(path, InputRole.R1)

    assert result.is_accepted
    assert result.records[0].campaign_member_ids == ("00vABCDEF123456XYZ",)
    assert result.records[1].campaign_member_ids == ("00v123456789ABCDEF",)
    assert result.records[2].campaign_member_ids == ("701AbCdEfGhIjKlIVK",)
    assert result.records[3].campaign_member_ids == ()
    assert IssueCode.CAMPAIGN_MEMBER_ID_INVALID in {
        issue.code for issue in result.issues
    }
    source_cell = result.source.records[1].cell("Campaigm Member")  # type: ignore[union-attr]
    assert source_cell.hyperlink_target is not None


def test_blank_r3_ids_are_allowed_and_distinct_rows_remain_countable(tmp_path):
    path = tmp_path / "r3.xlsx"
    _write_workbook(
        path,
        InputRole.R3,
        [
            _row(InputRole.R3, **{"First Name": "Один"}),
            _row(InputRole.R3, **{"First Name": "Два"}),
        ],
    )

    result = read_workbook(path, InputRole.R3)

    assert result.is_accepted
    assert not result.issues
    assert len(result.countable_records) == 2
    assert result.record_accounting.is_complete
    assert {group.status for group in result.record_groups} == {
        RecordGroupStatus.UNIQUE
    }


def test_identical_id_is_counted_once_but_all_sources_are_preserved(tmp_path):
    path = tmp_path / "r4.xlsx"
    row = _row(
        InputRole.R4,
        **{"Activity ID": "00T000000000001", "Name": "Учасник"},
    )
    _write_workbook(path, InputRole.R4, [row, row])

    result = read_workbook(path, InputRole.R4)

    assert result.source is not None and result.source.row_count == 2
    assert len(result.records) == 2
    assert len(result.countable_records) == 1
    assert not result.uncertain_groups
    assert result.record_groups[0].status is RecordGroupStatus.IDENTICAL_ID
    assert result.issues[0].code is IssueCode.DUPLICATE_ID
    assert len(result.issues[0].sources) == 2


def test_conflicting_id_and_identical_rows_without_id_have_no_winner(tmp_path):
    conflicting = tmp_path / "conflicting.xlsx"
    _write_workbook(
        conflicting,
        InputRole.R4,
        [
            _row(
                InputRole.R4,
                **{"Activity ID": "00T000000000001", "Name": "Один"},
            ),
            _row(
                InputRole.R4,
                **{"Activity ID": "00T000000000001", "Name": "Інший"},
            ),
        ],
    )
    conflict_result = read_workbook(conflicting, InputRole.R4)
    assert not conflict_result.countable_records
    assert len(conflict_result.uncertain_groups) == 1
    assert not conflict_result.record_accounting.is_complete
    assert conflict_result.record_groups[0].status is RecordGroupStatus.CONFLICTING_ID
    assert conflict_result.issues[0].code is IssueCode.CONFLICTING_ID_DATA
    assert "Name" in dict(conflict_result.issues[0].details)["different_fields"]

    uncertain = tmp_path / "uncertain.xlsx"
    row = _row(InputRole.R3, **{"First Name": "Тезка", "Last Name": "Один"})
    _write_workbook(uncertain, InputRole.R3, [row, row])
    uncertain_result = read_workbook(uncertain, InputRole.R3)
    assert not uncertain_result.countable_records
    assert len(uncertain_result.uncertain_groups) == 1
    assert uncertain_result.record_groups[0].status is RecordGroupStatus.UNCERTAIN_NO_ID
    assert uncertain_result.issues[0].code is IssueCode.DUPLICATE_UNCERTAIN


def test_numeric_identifier_and_missing_quantity_are_record_issues(tmp_path):
    path = tmp_path / "values.xlsx"
    _write_workbook(
        path,
        InputRole.R2,
        [
            _row(
                InputRole.R2,
                **{
                    "Opportunity Product Id 18": 123,
                    "Tax ID 1": 1234,
                    "Voucher Number": "SE-001",
                },
            ),
            _row(
                InputRole.R2,
                **{
                    "Opportunity Product Id 18": "000123",
                    "Tax ID 1": "001234",
                    "Voucher Number": "SE-002",
                    "Current Year Planned Quantity": 0,
                },
            ),
        ],
    )

    result = read_workbook(path, InputRole.R2)

    assert result.is_accepted
    assert result.records[0].product_id == "123"
    assert result.records[0].tax_id == "1234"
    assert result.records[1].product_id == "000123"
    assert result.records[1].tax_id == "001234"
    assert str(result.records[1].quantity) == "0"
    assert [issue.code for issue in result.issues].count(
        IssueCode.NUMERIC_ID_RISK
    ) == 2
    assert IssueCode.QUANTITY_UNKNOWN in {issue.code for issue in result.issues}


def test_conflicting_product_id_marks_its_contribution_unknown(tmp_path):
    path = tmp_path / "product-conflict.xlsx"
    base = {
        "Opportunity Product Id 18": "00k000000000001",
        "Tax ID 1": "001234",
        "Voucher Number": "SE-001",
    }
    _write_workbook(
        path,
        InputRole.R2,
        [
            _row(
                InputRole.R2,
                **{**base, "Current Year Planned Quantity": 1},
            ),
            _row(
                InputRole.R2,
                **{**base, "Current Year Planned Quantity": 2},
            ),
        ],
    )

    result = read_workbook(path, InputRole.R2)

    assert not result.countable_records
    assert len(result.uncertain_groups) == 1
    assert result.record_groups[0].status is RecordGroupStatus.CONFLICTING_ID
    conflict = next(
        issue
        for issue in result.issues
        if issue.code is IssueCode.PRODUCT_ID_CONFLICT
    )
    assert "Current Year Planned Quantity" in dict(conflict.details)[
        "different_fields"
    ]


def test_record_accounting_keeps_known_records_and_uncertain_groups(tmp_path):
    path = tmp_path / "participant-conflict.xlsx"
    shared_id = "701AbCdEfGhIjKlIVK"
    _write_workbook(
        path,
        InputRole.R3,
        [
            _row(
                InputRole.R3,
                **{
                    "Campaign Member Id 18": shared_id,
                    "First Name": "Один",
                },
            ),
            _row(
                InputRole.R3,
                **{
                    "Campaign Member Id 18": shared_id,
                    "First Name": "Суперечливий",
                },
            ),
            _row(InputRole.R3, **{"First Name": "Відомий"}),
        ],
    )

    result = read_workbook(path, InputRole.R3)

    accounting = result.record_accounting
    assert len(accounting.countable_records) == 1
    assert len(accounting.uncertain_groups) == 1
    assert accounting.uncertain_groups[0].status is RecordGroupStatus.CONFLICTING_ID
    assert not accounting.is_complete


def test_internal_grouping_failure_is_not_reported_as_invalid_xlsx(
    tmp_path, monkeypatch
):
    path = tmp_path / "internal.xlsx"
    _write_workbook(
        path,
        InputRole.R3,
        [_row(InputRole.R3, **{"First Name": "Учасник"})],
    )

    def fail_grouping(*_args, **_kwargs):
        raise KeyError("programmer defect")

    logged: list[str] = []
    monkeypatch.setattr(workbook_reader, "_group_records", fail_grouping)
    monkeypatch.setattr(
        workbook_reader._LOGGER,
        "error",
        lambda message, *args: logged.append(message % args),
    )
    result = read_workbook(path, InputRole.R3)

    assert not result.is_accepted
    assert result.source is not None
    assert not result.blocking_issues
    assert [issue.code for issue in result.fatal_issues] == [
        IssueCode.INTERNAL_ERROR
    ]
    issue = result.fatal_issues[0]
    assert issue.level.value == "internal"
    assert dict(issue.details)["phase"] == "record_grouping"
    assert "пошкодж" not in issue.message_uk
    assert "Стек викликів без значень джерела" in logged[0]
    assert dict(issue.details)["diagnostic_ref"] in logged[0]

    import_result = import_workbooks(_empty_set(tmp_path))
    assert import_result.snapshot is None
    assert not import_result.is_accepted
    assert import_result.fatal_issues
    assert all(
        issue.code is IssueCode.INTERNAL_ERROR
        for issue in import_result.fatal_issues
    )


def test_salesforce_15_and_18_ids_share_one_comparison_key(tmp_path):
    short = "701AbCdEfGhIjKl"
    full = "701AbCdEfGhIjKlIVK"

    assert salesforce_id_18(short) == full
    assert salesforce_id_key(short) == salesforce_id_key(full)
    assert salesforce_id_key(full.swapcase()) == salesforce_id_key(full)

    path = tmp_path / "r3-short-id.xlsx"
    _write_workbook(
        path,
        InputRole.R3,
        [_row(InputRole.R3, **{"Campaign Member Id 18": short})],
    )
    result = read_workbook(path, InputRole.R3)
    assert result.records[0].campaign_member_id == full


def test_missing_slot_prevents_snapshot_but_other_slots_are_still_checked(tmp_path):
    paths = _empty_set(tmp_path)
    del paths[InputRole.R2]

    result = import_workbooks(paths)

    assert not result.is_accepted
    assert result.snapshot is None
    assert result.workbook(InputRole.R2).issues[0].code is IssueCode.MISSING_INPUT
    assert all(
        result.workbook(role).is_accepted
        for role in (InputRole.R1, InputRole.R3, InputRole.R4)
    )
