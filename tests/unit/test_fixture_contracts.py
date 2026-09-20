from __future__ import annotations

from decimal import Decimal
from hashlib import sha256
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from seedlink.domain.provenance import InputRole
from seedlink.input_xlsx.schemas import schema_for


PROJECT_ROOT = Path(__file__).resolve().parents[2]
FIXTURES = PROJECT_ROOT / "tests" / "fixtures"


def _load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def test_real_fixture_manifest_is_independent_and_matches_schemas():
    manifest = _load_json(FIXTURES / "real" / "manifest.json")
    assert manifest["fixture_version"] == 1
    assert {source["role"] for source in manifest["sources"]} == {
        role.value for role in InputRole
    }
    for source in manifest["sources"]:
        assert tuple(source["columns"]) == schema_for(InputRole(source["role"])).columns

    expected = manifest["expected"]
    assert Decimal(expected["sunflower_quantity"]) + Decimal(
        expected["maize_quantity"]
    ) == Decimal(expected["known_quantity"])
    shared = expected["matched_vouchers"][2]
    assert sum(map(Decimal, shared["client_quantity_parts"])) == Decimal(
        shared["quantity"]
    )


def test_local_real_workbooks_match_recorded_fingerprints_and_structure():
    manifest = _load_json(FIXTURES / "real" / "manifest.json")
    missing = [
        source["path_from_project"]
        for source in manifest["sources"]
        if not (PROJECT_ROOT / source["path_from_project"]).is_file()
    ]
    if missing:
        pytest.skip(f"private local fixture is absent: {missing}")

    workbooks = {}
    for source in manifest["sources"]:
        path = PROJECT_ROOT / source["path_from_project"]
        assert sha256(path.read_bytes()).hexdigest() == source["sha256"]
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            assert len(workbook.worksheets) == source["sheet_count"]
            sheet = workbook.worksheets[0]
            assert sheet.title == source["sheet_name"]
            headers = tuple(cell.value for cell in next(sheet.iter_rows(max_row=1)))
            assert headers == tuple(source["columns"])
            assert sheet.max_row - 1 == source["data_rows"]
            workbooks[source["role"]] = [
                dict(zip(headers, (cell.value for cell in row), strict=True))
                for row in sheet.iter_rows(min_row=2)
            ]
        finally:
            workbook.close()

    expected = manifest["expected"]
    member_types = {}
    for row in workbooks["R3"]:
        member_type = row["Member Type"]
        member_types[member_type] = member_types.get(member_type, 0) + 1
    assert member_types == expected["member_types"]
    assert sum(bool(row["Campaign Member Id 18"]) for row in workbooks["R3"]) == 0

    r2_by_excel_row = {
        excel_row: row for excel_row, row in enumerate(workbooks["R2"], start=2)
    }
    species_totals = {"sunflower": Decimal(0), "maize": Decimal(0)}
    tax_totals = {}
    for voucher in expected["matched_vouchers"]:
        rows = [r2_by_excel_row[row_number] for row_number in voucher["r2_excel_rows"]]
        assert {row["Voucher Number"] for row in rows} == {voucher["full_number"]}
        assert sum(Decimal(row["Current Year Planned Quantity"]) for row in rows) == Decimal(
            voucher["quantity"]
        )
        for row in rows:
            quantity = Decimal(row["Current Year Planned Quantity"])
            species = row["Species group"]
            if species == "SUN: Sunflowers":
                species_totals["sunflower"] += quantity
            elif species == "CRN: Corn Seeds":
                species_totals["maize"] += quantity
            tax_id = row["Tax ID 1"]
            tax_totals[tax_id] = tax_totals.get(tax_id, Decimal(0)) + quantity
    assert species_totals == {
        "sunflower": Decimal(expected["sunflower_quantity"]),
        "maize": Decimal(expected["maize_quantity"]),
    }
    assert sorted(tax_totals.values()) == [Decimal("7"), Decimal("33"), Decimal("66")]


def test_synthetic_fixture_is_anonymous_and_has_hand_checked_totals():
    fixture = _load_json(FIXTURES / "synthetic" / "complex_relationships.json")
    serialized = json.dumps(fixture, ensure_ascii=False)
    assert "2777788899" not in serialized
    assert "@" not in serialized

    accepted = set(fixture["expected"]["accepted_voucher_keys"])
    accepted_lines = [
        line for line in fixture["product_lines"] if line["voucher_key"] in accepted
    ]
    assert sum(Decimal(line["quantity"]) for line in accepted_lines) == Decimal(
        fixture["expected"]["known_accepted_quantity"]
    )
    assert fixture["expected"]["survey_voucher_links"]["S-002"]["accepted"] == []
    assert set(
        fixture["expected"]["survey_voucher_links"]["S-001"]["accepted"]
    ) == {"V-001", "V-002"}
